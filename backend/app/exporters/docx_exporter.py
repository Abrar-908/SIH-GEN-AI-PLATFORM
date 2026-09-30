import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.core.config import EXPORTS_DIR

class DocxExporter:
    @staticmethod
    def generate_docx(
        filename: str,
        title: str,
        content_markdown: str,
        output_type: str,
        source_references: list = None
    ) -> str:
        filepath = EXPORTS_DIR / filename
        doc = Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)

        # Title
        title_p = doc.add_paragraph()
        run_title = title_p.add_run(title)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        # Subtitle
        sub_p = doc.add_paragraph()
        run_sub = sub_p.add_run(f"IntelTransform AI — Source-Grounded {output_type}")
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(2, 132, 199)

        doc.add_paragraph("-" * 65)

        # Markdown body
        lines = content_markdown.split("\n")
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("# "):
                continue
            elif line_str.startswith("## ") or line_str.startswith("### "):
                clean_h = line_str.lstrip("#").strip()
                h_p = doc.add_paragraph()
                h_run = h_p.add_run(clean_h)
                h_run.font.name = "Calibri"
                h_run.font.size = Pt(14)
                h_run.font.bold = True
                h_run.font.color.rgb = RGBColor(30, 41, 59)
            elif line_str.startswith("* ") or line_str.startswith("- "):
                clean_bullet = line_str[2:].strip().replace("**", "")
                b_p = doc.add_paragraph(style='List Bullet')
                b_run = b_p.add_run(clean_bullet)
                b_run.font.name = "Calibri"
                b_run.font.size = Pt(10.5)
            else:
                p = doc.add_paragraph()
                p_run = p.add_run(line_str.replace("**", ""))
                p_run.font.name = "Calibri"
                p_run.font.size = Pt(10.5)
                p_run.font.color.rgb = RGBColor(51, 65, 85)

        # Citations
        if source_references:
            doc.add_paragraph("-" * 65)
            h_ref = doc.add_paragraph()
            h_ref_run = h_ref.add_run("Grounding Citations & Source Audit")
            h_ref_run.font.name = "Calibri"
            h_ref_run.font.size = Pt(13)
            h_ref_run.font.bold = True

            table = doc.add_table(rows=1, cols=4)
            hdr_cells = table.rows[0].cells
            headers = ["Claim", "Source Page", "Chunk ID", "Confidence"]
            for i, h_text in enumerate(headers):
                hdr_cells[i].text = h_text

            for ref in source_references[:6]:
                row_cells = table.add_row().cells
                row_cells[0].text = ref.get("claim", "")[:80]
                row_cells[1].text = f"Page {ref.get('source_page', 1)}"
                row_cells[2].text = ref.get("source_chunk", "chunk_01")
                row_cells[3].text = f"{int(ref.get('confidence', 0.95) * 100)}%"

        doc.save(str(filepath))
        return str(filepath)
