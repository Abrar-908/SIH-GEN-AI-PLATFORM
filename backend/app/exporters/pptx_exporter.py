import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from app.core.config import EXPORTS_DIR

class PPTXExporter:
    @staticmethod
    def generate_pptx(
        filename: str,
        title: str,
        structured_slides: list = None,
        content_markdown: str = ""
    ) -> str:
        filepath = EXPORTS_DIR / filename
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 widescreen
        prs.slide_height = Inches(7.5)

        # Blank slide layout
        blank_layout = prs.slide_layouts[6]

        # Dark theme background color
        BG_COLOR = RGBColor(11, 15, 25)
        ACCENT_CYAN = RGBColor(6, 182, 212)
        TEXT_WHITE = RGBColor(241, 245, 249)
        TEXT_MUTED = RGBColor(148, 163, 184)
        CARD_BG = RGBColor(17, 24, 39)

        if not structured_slides:
            # Fallback slide structure
            structured_slides = [
                {
                    "slide_number": 1,
                    "title": title,
                    "subtitle": "IntelTransform AI — Source-Grounded Executive Briefing",
                    "bullets": ["Automated multi-format transformation", "Verified citations", "Executive-ready synthesis"],
                    "notes": "Introduction slide generated automatically by IntelTransform AI."
                }
            ]

        for s_idx, slide_data in enumerate(structured_slides):
            slide = prs.slides.add_slide(blank_layout)

            # Slide background shape
            bg_shape = slide.shapes.add_shape(
                1, 0, 0, Inches(13.333), Inches(7.5) # 1 is msoShapeRectangle
            )
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = BG_COLOR
            bg_shape.line.color.rgb = BG_COLOR

            # Top branding header
            brand_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
            brand_tf = brand_box.text_frame
            brand_p = brand_tf.paragraphs[0]
            brand_p.text = "INTELTRANSFORM AI  |  SECURE CONTENT SYNTHESIS PLATFORM"
            brand_p.font.name = "Arial"
            brand_p.font.size = Pt(10)
            brand_p.font.bold = True
            brand_p.font.color.rgb = ACCENT_CYAN

            # Slide Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(1.0))
            title_tf = title_box.text_frame
            title_tf.word_wrap = True
            title_p = title_tf.paragraphs[0]
            title_p.text = slide_data.get("title", f"Slide {s_idx + 1}")
            title_p.font.name = "Arial"
            title_p.font.size = Pt(28)
            title_p.font.bold = True
            title_p.font.color.rgb = TEXT_WHITE

            # Subtitle
            subtitle_text = slide_data.get("subtitle", "")
            if subtitle_text:
                sub_p = title_tf.add_paragraph()
                sub_p.text = subtitle_text
                sub_p.font.name = "Arial"
                sub_p.font.size = Pt(14)
                sub_p.font.color.rgb = ACCENT_CYAN

            # Content Card
            card_shape = slide.shapes.add_shape(
                1, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5)
            )
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = CARD_BG
            card_shape.line.color.rgb = RGBColor(30, 41, 59)

            content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.9), Inches(3.9))
            content_tf = content_box.text_frame
            content_tf.word_wrap = True

            bullets = slide_data.get("bullets", [])
            for b_idx, bullet in enumerate(bullets):
                p = content_tf.paragraphs[0] if b_idx == 0 else content_tf.add_paragraph()
                p.text = f"•  {bullet}"
                p.font.name = "Arial"
                p.font.size = Pt(16)
                p.font.color.rgb = TEXT_WHITE
                p.space_after = Pt(16)

            # Slide Notes
            notes = slide_data.get("notes", "")
            if notes and slide.notes_slide:
                notes_tf = slide.notes_slide.notes_text_frame
                notes_tf.text = notes

        prs.save(str(filepath))
        return str(filepath)
